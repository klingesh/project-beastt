"""Turn a structured spec into a real Office document.

The LLM produces a JSON spec (see docgen.py); this module renders it into an
actual .pptx / .docx / .xlsx file. Rendering is kept separate from generation so
the layout code is deterministic and testable, and a weak model can't produce a
corrupt file -- only a thin or oddly-worded one.

Each builder degrades gracefully: if the relevant library isn't installed, it
raises a clear message telling the user exactly what to install.
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Dict, List

from .paths import project_root


def output_dir() -> Path:
    path = project_root() / "beastt_output"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _safe_name(title: str, suffix: str) -> Path:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", title or "document").strip("_")[:60]
    slug = slug or "document"
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    return output_dir() / f"{slug}_{stamp}{suffix}"


class MissingLibrary(RuntimeError):
    pass


# --- PowerPoint -------------------------------------------------------------
def build_presentation(spec: Dict, theme_name: str = "navy") -> Path:
    """Render a designed 16:9 slide deck.

    Spec: {title, subtitle, slides:[{title, bullets[], notes, key_message}]}

    Rather than using stock Office layouts (which look plainly templated), each
    slide is composed from scratch: a coloured title band, an accent rule, a
    typographic hierarchy, and a footer with slide numbers.
    """
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Emu, Inches, Pt
    except Exception as exc:
        raise MissingLibrary(
            f"PowerPoint support needs python-pptx ({exc}). Run: pip install python-pptx"
        )

    from .theme import get as get_theme

    th = get_theme(theme_name)
    title = spec.get("title") or "Presentation"
    subtitle = spec.get("subtitle") or datetime.now().strftime("%B %d, %Y")

    prs = Presentation()
    prs.slide_width = Inches(13.333)   # 16:9 widescreen
    prs.slide_height = Inches(7.5)
    SW, SH = prs.slide_width, prs.slide_height
    BLANK = prs.slide_layouts[6]

    def rect(slide, left, top, width, height, colour):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = th.rgb(colour)
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def textbox(slide, left, top, width, height):
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.word_wrap = True
        return frame

    def write(frame, text, size, colour, bold=False, font=None, align=PP_ALIGN.LEFT,
              space_after=6, first=False):
        para = frame.paragraphs[0] if first else frame.add_paragraph()
        para.alignment = align
        para.space_after = Pt(space_after)
        run = para.add_run()
        run.text = str(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font or th.body_font
        run.font.color.rgb = th.rgb(colour)
        return para

    # ---- Title slide: full-bleed colour with a large accent rule ----------
    slide = prs.slides.add_slide(BLANK)
    rect(slide, 0, 0, SW, SH, th.primary)
    rect(slide, Inches(0.9), Inches(2.55), Inches(1.6), Pt(6), th.accent)

    frame = textbox(slide, Inches(0.9), Inches(2.8), SW - Inches(2.4), Inches(2.2))
    write(frame, title, th.title_size, th.white, bold=True, font=th.heading_font, first=True)
    write(frame, subtitle, th.subtitle_size, th.accent, space_after=0)

    frame = textbox(slide, Inches(0.9), SH - Inches(1.0), SW - Inches(2), Inches(0.4))
    write(frame, datetime.now().strftime("%B %d, %Y"), th.caption_size,
          th.light, first=True)

    slides = [s for s in spec.get("slides", []) if isinstance(s, dict)]

    # ---- Agenda slide, derived from the deck's own structure --------------
    headings = [str(s.get("title") or "").strip() for s in slides]
    headings = [h for h in headings if h]
    if len(headings) >= 3:
        slide = prs.slides.add_slide(BLANK)
        rect(slide, 0, 0, SW, Inches(1.15), th.primary)
        frame = textbox(slide, Inches(0.7), Inches(0.25), SW - Inches(1.4), Inches(0.7))
        write(frame, "Agenda", th.slide_title_size, th.white, bold=True,
              font=th.heading_font, first=True)

        frame = textbox(slide, Inches(0.9), Inches(1.7), SW - Inches(1.8), SH - Inches(2.6))
        for idx, heading in enumerate(headings, 1):
            para = write(frame, f"{idx:02d}    {heading}", 20, th.text_dark,
                         space_after=14, first=(idx == 1))
            para.runs[0].font.name = th.body_font

    # ---- Content slides ---------------------------------------------------
    total = len(slides)
    for number, item in enumerate(slides, 1):
        slide = prs.slides.add_slide(BLANK)

        # Header band + accent rule
        rect(slide, 0, 0, SW, Inches(1.15), th.primary)
        rect(slide, 0, Inches(1.15), SW, Pt(4), th.accent)

        frame = textbox(slide, Inches(0.7), Inches(0.22), SW - Inches(1.4), Inches(0.8))
        write(frame, str(item.get("title") or ""), th.slide_title_size, th.white,
              bold=True, font=th.heading_font, first=True)

        bullets = [str(b).strip() for b in (item.get("bullets") or []) if str(b).strip()]
        key = str(item.get("key_message") or "").strip()

        body_top = Inches(1.65)
        body_height = SH - body_top - Inches(1.2)

        if bullets:
            # Two columns once there are enough bullets to look sparse in one.
            two_col = len(bullets) >= 5
            col_width = (SW - Inches(1.8)) / (2 if two_col else 1)
            groups = (
                [bullets[: (len(bullets) + 1) // 2], bullets[(len(bullets) + 1) // 2 :]]
                if two_col
                else [bullets]
            )
            for col, group in enumerate(groups):
                if not group:
                    continue
                left = Inches(0.9) + Emu(int(col_width)) * col
                frame = textbox(slide, left, body_top,
                                Emu(int(col_width)) - Inches(0.3), body_height)
                frame.vertical_anchor = MSO_ANCHOR.TOP
                for idx, bullet in enumerate(group):
                    para = write(frame, f"▪   {bullet}", th.bullet_size, th.text_dark,
                                 space_after=12, first=(idx == 0))
                    para.line_spacing = 1.15

        # Key takeaway panel, when the model supplied one.
        if key:
            panel_top = SH - Inches(1.55)
            rect(slide, Inches(0.9), panel_top, SW - Inches(1.8), Inches(0.72), th.light)
            rect(slide, Inches(0.9), panel_top, Pt(5), Inches(0.72), th.accent)
            frame = textbox(slide, Inches(1.15), panel_top + Inches(0.12),
                            SW - Inches(2.3), Inches(0.5))
            write(frame, key, 14, th.secondary, bold=True, first=True)

        # Footer: deck title and slide number.
        frame = textbox(slide, Inches(0.9), SH - Inches(0.55), SW - Inches(3), Inches(0.35))
        write(frame, title, th.caption_size, th.text_muted, first=True)
        frame = textbox(slide, SW - Inches(1.6), SH - Inches(0.55), Inches(0.9), Inches(0.35))
        write(frame, f"{number} / {total}", th.caption_size, th.text_muted,
              align=PP_ALIGN.RIGHT, first=True)

        notes = item.get("notes")
        if notes:
            slide.notes_slide.notes_text_frame.text = str(notes)

    # ---- Closing slide ----------------------------------------------------
    slide = prs.slides.add_slide(BLANK)
    rect(slide, 0, 0, SW, SH, th.primary)
    rect(slide, Inches(0.9), Inches(3.15), Inches(1.6), Pt(6), th.accent)
    frame = textbox(slide, Inches(0.9), Inches(3.4), SW - Inches(2), Inches(1.2))
    write(frame, spec.get("closing") or "Thank you", 36, th.white, bold=True,
          font=th.heading_font, first=True)

    path = _safe_name(title, ".pptx")
    prs.save(path)
    return path


# --- Word -------------------------------------------------------------------
def build_document(spec: Dict, theme_name: str = "navy") -> Path:
    """Render a styled report.

    Spec: {title, subtitle, sections:[{heading, paragraphs[], bullets[]}]}

    Adds a cover page, themed headings, justified body text, and a footer, so the
    result reads as a report rather than raw text in the default template.
    """
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor
    except Exception as exc:
        raise MissingLibrary(
            f"Word support needs python-docx ({exc}). Run: pip install python-docx"
        )

    from .theme import get as get_theme

    th = get_theme(theme_name)
    title = spec.get("title") or "Document"
    subtitle = spec.get("subtitle") or ""

    doc = Document()

    # Base body style.
    normal = doc.styles["Normal"]
    normal.font.name = th.body_font
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    def colour(hex_colour):
        return RGBColor.from_string(hex_colour)

    # Themed heading styles.
    for style_name, size, hexc in (
        ("Title", 30, th.primary),
        ("Heading 1", 17, th.primary),
        ("Heading 2", 13, th.secondary),
    ):
        try:
            style = doc.styles[style_name]
            style.font.name = th.heading_font
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = colour(hexc)
        except KeyError:
            pass

    # ---- Cover ----
    heading = doc.add_paragraph(title, style="Title")
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if subtitle:
        para = doc.add_paragraph()
        run = para.add_run(subtitle)
        run.font.size = Pt(13)
        run.font.color.rgb = colour(th.text_muted)
        run.italic = True

    para = doc.add_paragraph()
    run = para.add_run(datetime.now().strftime("%B %d, %Y"))
    run.font.size = Pt(10)
    run.font.color.rgb = colour(th.text_muted)

    sections = [s for s in spec.get("sections", []) if isinstance(s, dict)]

    # ---- Contents, for anything substantial ----
    headings = [str(s.get("heading") or "").strip() for s in sections]
    headings = [h for h in headings if h]
    if len(headings) >= 3:
        doc.add_paragraph("Contents", style="Heading 1")
        for idx, item in enumerate(headings, 1):
            para = doc.add_paragraph(f"{idx}.  {item}")
            para.paragraph_format.space_after = Pt(2)

    # ---- Body ----
    for section in sections:
        head = section.get("heading")
        if head:
            doc.add_paragraph(str(head), style="Heading 1")
        for text in section.get("paragraphs", []) or []:
            if str(text).strip():
                para = doc.add_paragraph(str(text))
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for bullet in section.get("bullets", []) or []:
            if str(bullet).strip():
                doc.add_paragraph(str(bullet), style="List Bullet")

    # ---- Footer ----
    try:
        footer = doc.sections[0].footer.paragraphs[0]
        footer.text = title
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = colour(th.text_muted)
    except Exception:
        pass

    path = _safe_name(title, ".docx")
    doc.save(path)
    return path


# --- Excel ------------------------------------------------------------------
def build_spreadsheet(spec: Dict, theme_name: str = "navy") -> Path:
    """Render a formatted workbook.

    Spec: {title, sheets:[{name, columns[], rows[[...]]}]}

    Applies a themed header row, frozen panes, autofilter, banded rows, and
    number formatting so it's usable immediately rather than a raw dump.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except Exception as exc:
        raise MissingLibrary(
            f"Excel support needs openpyxl ({exc}). Run: pip install openpyxl"
        )

    from .theme import get as get_theme

    th = get_theme(theme_name)
    title = spec.get("title") or "Workbook"
    wb = Workbook()
    sheets = spec.get("sheets") or [{"name": "Sheet1", "columns": [], "rows": []}]

    header_fill = PatternFill("solid", fgColor=th.primary)
    band_fill = PatternFill("solid", fgColor=th.light)
    header_font = Font(bold=True, color=th.white, name=th.body_font, size=11)
    body_font = Font(name=th.body_font, size=11)
    thin = Side(style="thin", color="D6DEE7")
    border = Border(bottom=thin)

    for idx, sheet in enumerate(sheets):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = str(sheet.get("name") or f"Sheet{idx + 1}")[:31]

        columns = [str(c) for c in (sheet.get("columns") or [])]
        if columns:
            ws.append(columns)
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center",
                                           wrap_text=True)
            ws.row_dimensions[1].height = 22
            ws.freeze_panes = "A2"

        for row in sheet.get("rows") or []:
            if isinstance(row, dict):  # tolerate dict rows keyed by column
                ws.append([row.get(col, "") for col in columns])
            elif isinstance(row, (list, tuple)):
                ws.append(list(row))
            else:
                ws.append([row])

        # Body styling: banded rows, thousands separators, tidy alignment.
        for r, row in enumerate(ws.iter_rows(min_row=2), start=2):
            for cell in row:
                cell.font = body_font
                cell.border = border
                if r % 2 == 0:
                    cell.fill = band_fill
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cell.number_format = "#,##0.##"
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.alignment = Alignment(vertical="center", wrap_text=False)

        if columns and ws.max_row > 1:
            ws.auto_filter.ref = (
                f"A1:{get_column_letter(len(columns))}{ws.max_row}"
            )

        # Size columns to their content.
        widths: List[int] = []
        for row in ws.iter_rows(values_only=True):
            for i, value in enumerate(row):
                length = len(str(value)) if value is not None else 0
                if i >= len(widths):
                    widths.append(length)
                else:
                    widths[i] = max(widths[i], length)
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 4, 12), 55)

    path = _safe_name(title, ".xlsx")
    wb.save(path)
    return path


BUILDERS = {
    "presentation": build_presentation,
    "document": build_document,
    "spreadsheet": build_spreadsheet,
}


def build(kind: str, spec: Dict, theme_name: str = "navy") -> Path:
    builder = BUILDERS.get(kind)
    if builder is None:
        raise ValueError(f"Unknown document kind: {kind}")
    return builder(spec, theme_name)
