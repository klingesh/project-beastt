"""Read documents: PDF, Word, Excel, PowerPoint, and plain text.

Each reader returns plain text plus a short note about the structure (pages,
sheets, slides), so a summary can say where something came from. Readers are
chosen by extension and degrade with a clear message when the relevant library
isn't installed.

Text is truncated to a sane size before it reaches the model -- a 200-page PDF
would otherwise blow past the context window and produce nothing useful.
"""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Optional, Tuple

MAX_CHARS = 12_000


class Unsupported(RuntimeError):
    pass


def _truncate(text: str, limit: int = MAX_CHARS) -> Tuple[str, bool]:
    text = text.strip()
    if len(text) <= limit:
        return text, False
    return text[:limit].rsplit(" ", 1)[0], True


# --- individual formats -----------------------------------------------------
def read_pdf(data: bytes) -> Tuple[str, str]:
    try:
        from pypdf import PdfReader
    except Exception:
        try:
            from PyPDF2 import PdfReader      # older installs
        except Exception as exc:
            raise Unsupported(
                f"PDF reading needs pypdf ({exc}). Run: pip install pypdf"
            )
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    body = "\n\n".join(p.strip() for p in pages if p.strip())
    if not body:
        raise Unsupported(
            "That PDF has no extractable text -- it's probably a scan, which would "
            "need OCR."
        )
    return body, f"{len(reader.pages)} page(s)"


def read_docx(data: bytes) -> Tuple[str, str]:
    try:
        from docx import Document
    except Exception as exc:
        raise Unsupported(
            f"Word reading needs python-docx ({exc}). Run: pip install python-docx"
        )
    doc = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts), f"{len(doc.paragraphs)} paragraph(s), {len(doc.tables)} table(s)"


def read_pptx(data: bytes) -> Tuple[str, str]:
    try:
        from pptx import Presentation
    except Exception as exc:
        raise Unsupported(
            f"PowerPoint reading needs python-pptx ({exc}). Run: pip install python-pptx"
        )
    prs = Presentation(io.BytesIO(data))
    blocks = []
    for index, slide in enumerate(prs.slides, 1):
        lines = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if lines:
            blocks.append(f"[Slide {index}]\n" + "\n".join(lines))
    return "\n\n".join(blocks), f"{len(prs.slides)} slide(s)"


def read_xlsx(data: bytes) -> Tuple[str, str]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise Unsupported(
            f"Excel reading needs openpyxl ({exc}). Run: pip install openpyxl"
        )
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    blocks = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells).rstrip(" |"))
            if len(rows) >= 200:      # enough to characterise a sheet
                rows.append("...")
                break
        if rows:
            blocks.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
    return "\n\n".join(blocks), f"{len(workbook.worksheets)} sheet(s)"


def read_csv(data: bytes) -> Tuple[str, str]:
    text = data.decode("utf-8", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    lines = [" | ".join(r) for r in rows[:200]]
    return "\n".join(lines), f"{len(rows)} row(s)"


def read_text(data: bytes) -> Tuple[str, str]:
    text = data.decode("utf-8", errors="replace")
    return text, f"{len(text.splitlines())} line(s)"


def read_json(data: bytes) -> Tuple[str, str]:
    text = data.decode("utf-8", errors="replace")
    try:
        parsed = json.loads(text)
        pretty = json.dumps(parsed, indent=1)[:MAX_CHARS]
        shape = (
            f"{len(parsed)} key(s)" if isinstance(parsed, dict)
            else f"{len(parsed)} item(s)" if isinstance(parsed, list) else "value"
        )
        return pretty, shape
    except Exception:
        return text, "invalid JSON, read as text"


READERS = {
    ".pdf": read_pdf,
    ".docx": read_docx, ".doc": read_docx,
    ".pptx": read_pptx, ".ppt": read_pptx,
    ".xlsx": read_xlsx, ".xlsm": read_xlsx,
    ".csv": read_csv,
    ".json": read_json,
    ".txt": read_text, ".md": read_text, ".rst": read_text, ".log": read_text,
    ".py": read_text, ".js": read_text, ".ts": read_text, ".html": read_text,
    ".css": read_text, ".yml": read_text, ".yaml": read_text, ".sql": read_text,
    ".ini": read_text, ".cfg": read_text, ".toml": read_text,
}

DOCUMENT_SUFFIXES = tuple(READERS)


def supported(name: str) -> bool:
    return Path(name).suffix.lower() in READERS


def extract(name: str, data: bytes) -> Tuple[str, str, bool]:
    """Return (text, structure_note, was_truncated) for a document."""
    suffix = Path(name).suffix.lower()
    reader = READERS.get(suffix)
    if reader is None:
        raise Unsupported(
            f"I can't read '{suffix or 'that'}' files. I handle PDF, Word, Excel, "
            "PowerPoint, CSV, JSON and plain text."
        )
    if suffix in (".doc", ".ppt"):
        # The old binary formats aren't readable by the modern libraries.
        raise Unsupported(
            f"'{suffix}' is the old binary format. Save it as "
            f"'{suffix}x' and I'll read it."
        )
    text, note = reader(data)
    text, truncated = _truncate(text)
    if not text.strip():
        raise Unsupported("There's no readable text in that file.")
    return text, note, truncated
